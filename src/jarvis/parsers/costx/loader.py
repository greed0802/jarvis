"""CostX workbook loader.

This module provides a workbook loader with CostX compatibility fixes.
Based on M5 engineering discovery findings for named style handling.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.workbook.workbook import Workbook
from openpyxl.styles.named_styles import _NamedCellStyle


# Track whether the patch has been applied
_NAMED_STYLE_PATCH_APPLIED = False


def _apply_named_style_none_name_patch() -> None:
    """Apply CostX compatibility patch for named styles with None names.

    CostX exports can contain cellXfs / named style entries with a missing name.
    openpyxl 3.1.x raises TypeError for these entries. This patch coerces None to ''.
    """
    global _NAMED_STYLE_PATCH_APPLIED
    if _NAMED_STYLE_PATCH_APPLIED:
        return

    # Apply the compatibility patch only once per process.
    original_init = _NamedCellStyle.__init__

    def patched_init(
        self,
        name: str | None = None,
        xfId: int | None = None,
        builtinId: int | None = None,
        iLevel: int | None = None,
        hidden: bool | None = None,
        customBuiltin: bool | None = None,
        extLst: Any = None,
    ) -> None:
        if name is None:
            name = ""
        original_init(
            self,
            name=name,
            xfId=xfId,
            builtinId=builtinId,
            iLevel=iLevel,
            hidden=hidden,
            customBuiltin=customBuiltin,
            extLst=extLst,
        )

    _NamedCellStyle.__init__ = patched_init
    _NAMED_STYLE_PATCH_APPLIED = True


def load_workbook(path: Path, *, data_only: bool = False) -> Workbook:
    """Load a workbook with CostX compatibility handling.

    Args:
        path: Path to the workbook file.
        data_only: If True, return calculated values instead of formulas.

    Returns:
        The loaded Workbook instance.

    Note:
        Exceptions from openpyxl.load_workbook() propagate if the workbook
        cannot be loaded.
    """
    try:
        return openpyxl.load_workbook(path, data_only=data_only)
    except TypeError as e:
        msg = str(e)
        if "NamedCellStyle" in msg and "NoneType" in msg:
            _apply_named_style_none_name_patch()
            return openpyxl.load_workbook(path, data_only=data_only)
        raise